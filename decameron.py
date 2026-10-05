"""DECAMERON: transcribes audio with the Whisper CLI and exports DOCX/PDF documents.

The raw transcript is cleaned up, split into sentences and regrouped into paragraphs
of 2-5 sentences, so the result reads like prose instead of one continuous block.
Each audio file gets its own folder (converted/<name>/) with the Whisper transcripts,
the documents and a README.txt recording when the conversion happened.
"""

import argparse
import importlib
import os
import random
import re
import shutil
import subprocess
import sys
import unicodedata
from collections import Counter
from collections.abc import Callable, Sequence
from datetime import datetime
from pathlib import Path
from typing import NamedTuple

RULE = "=" * 50
USE_COLOR = sys.stdout.isatty() and not os.environ.get("NO_COLOR")
PARAGRAPH_SIZE = (2, 5)  # min/max sentences per paragraph
OUTPUT_ROOT = Path("converted")
TIMESTAMP_FORMAT = "%Y-%m-%d %H:%M:%S"

# A sentence ends at . ! ? or an ellipsis, optionally followed by a closing quote or
# bracket, and then whitespace (the re module expands the escapes).
SENTENCE_END = re.compile(
    r"(?<=[.!?\u2026])\s+"
    r"|(?<=[.!?\u2026][\"'\u201d\u2019\u00bb)\]])\s+"
)

# fpdf2's core fonts (helvetica) only cover Windows-1252: Latin-1 plus curly quotes,
# dashes, "…", "€", etc. Any other character must be replaced before PDF export.
PDF_ENCODING = "windows-1252"
PDF_CHARSET = frozenset(bytes(range(256)).decode(PDF_ENCODING, errors="ignore"))
# Characters outside Windows-1252 that have no Unicode decomposition to fall back on
PDF_REPLACEMENTS = str.maketrans(
    {
        "\u2010": "-",  # hyphen
        "\u2011": "-",  # non-breaking hyphen
        "\u2212": "-",  # minus sign
        "\u2015": "\u2014",  # horizontal bar -> em dash
        "\u02bc": "\u2019",  # modifier letter apostrophe -> right single quote
        "\u2032": "'",  # prime
        "\u2033": '"',  # double prime
        "\u200b": None,  # zero-width space
        "\ufeff": None,  # byte order mark
    }
)


def bold(text: str) -> str:
    """Wraps text in ANSI bold codes when writing to a color-capable terminal."""
    return f"\033[1m{text}\033[0m" if USE_COLOR else text


def print_err(message: str = "") -> None:
    """Prints a message to stderr, flushing stdout first to keep the output in order."""
    sys.stdout.flush()
    print(message, file=sys.stderr)


def print_banner() -> None:
    """Prints the tool's title banner."""
    print(f"\n{bold('DECAMERON')}")
    print("(Audio Transcription & Formatting Tool)")
    print(RULE)


def split_sentences(text: str) -> list[str]:
    """Normalizes whitespace and splits the text into sentences."""
    sentences: list[str] = []
    for piece in SENTENCE_END.split(" ".join(text.split())):
        if sentences and sentences[-1].endswith(("...", "…")) and piece[:1].islower():
            # "Y entonces... no sé." is a pause, not a new sentence
            sentences[-1] += f" {piece}"
        elif piece:
            sentences.append(piece)
    return sentences


def format_text_to_paragraphs(text: str, rng: random.Random | None = None) -> list[str]:
    """Cleans up the text and groups its sentences into paragraphs of random size."""
    rng = rng or random.Random()
    low, high = PARAGRAPH_SIZE
    sentences = split_sentences(text)

    paragraphs: list[str] = []
    start = 0
    while start < len(sentences):
        size = rng.randint(low, high)
        if len(sentences) - start - size == 1:
            # Don't leave a lone sentence for the last paragraph
            size = size + 1 if size < high else size - 1
        paragraphs.append(" ".join(sentences[start : start + size]))
        start += size
    return paragraphs


def _approximate_for_pdf(char: str) -> str:
    """Approximates a character missing from the PDF font ('ő' -> 'o'), else '?'."""
    decomposed = unicodedata.normalize("NFKD", char)
    return "".join(c for c in decomposed if c in PDF_CHARSET) or "?"


def to_pdf_text(text: str) -> str:
    """Replaces the characters that fpdf2's core fonts cannot print."""
    text = unicodedata.normalize("NFC", text).translate(PDF_REPLACEMENTS)
    return "".join(c if c in PDF_CHARSET else _approximate_for_pdf(c) for c in text)


class DocumentContent(NamedTuple):
    """What goes into an exported document, in order."""

    title: str  # h1
    converted_at: str  # h3
    paragraphs: Sequence[str]


def save_as_docx(content: DocumentContent, path: Path) -> None:
    """Saves the content into a Microsoft Word (.docx) document."""
    from docx import Document

    doc = Document()
    doc.core_properties.title = content.title
    doc.add_heading(content.title, level=1)
    doc.add_paragraph()  # Line break
    doc.add_heading(content.converted_at, level=3)
    doc.add_paragraph()  # Line break
    for paragraph in content.paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))


def save_as_pdf(content: DocumentContent, path: Path) -> None:
    """Saves the content into a PDF document."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.core_fonts_encoding = PDF_ENCODING
    pdf.set_title(content.title)
    pdf.add_page()

    def write_block(text: str, size: int, line_height: float, style: str = "") -> None:
        pdf.set_font("helvetica", style=style, size=size)
        pdf.multi_cell(
            0, line_height, text=to_pdf_text(text), new_x="LMARGIN", new_y="NEXT"
        )

    write_block(content.title, size=22, line_height=10, style="B")  # h1
    pdf.ln(6)  # Line break
    write_block(content.converted_at, size=13, line_height=7, style="B")  # h3
    pdf.ln(6)  # Line break
    for paragraph in content.paragraphs:
        write_block(paragraph, size=11, line_height=6)
        pdf.ln(4)  # Extra space between paragraphs
    pdf.output(path)


def write_readme(
    folder: Path, audio: Path, converted_at: str, model: str, language: str
) -> Path:
    """Writes a README.txt describing the conversion and returns its path."""
    readme = folder / "README.txt"
    readme.write_text(
        "DECAMERON (Audio Transcription & Formatting Tool)\n\n"
        f"Converted on:  {converted_at}\n"
        f"Source file:   {audio.name}\n"
        f"Whisper model: {model}\n"
        f"Language:      {language}\n",
        encoding="utf-8",
    )
    return readme


class Exporter(NamedTuple):
    """An output format: the module it needs, its pip package and its save function."""

    module: str
    package: str
    save: Callable[[DocumentContent, Path], None]


EXPORTERS: dict[str, Exporter] = {
    "docx": Exporter("docx", "python-docx", save_as_docx),
    "pdf": Exporter("fpdf", "fpdf2", save_as_pdf),
}


def missing_dependencies(formats: Sequence[str]) -> dict[str, str]:
    """Maps each missing pip package needed for this run to the reason it is needed."""
    missing: dict[str, str] = {}
    for fmt in formats:
        exporter = EXPORTERS[fmt]
        try:
            importlib.import_module(exporter.module)
        except ImportError as e:  # not installed, or installed but broken
            missing[exporter.package] = f"needed for {fmt.upper()} output; {e}"
    if shutil.which("whisper") is None:
        missing["openai-whisper"] = "the 'whisper' command is not in your PATH"
    return missing


def run_whisper(
    audio: Path, model: str, language: str, output_dir: Path
) -> Path | None:
    """Runs the Whisper CLI on one file and returns its new .txt transcript."""
    print(f"\n🎙️ Transcribing with Whisper: {audio}...", flush=True)

    transcript = output_dir / f"{audio.stem}.txt"
    previous_mtime = transcript.stat().st_mtime_ns if transcript.exists() else None
    command = [
        "whisper",
        str(audio.absolute()),  # never mistaken for an option, even if named "-x.mp3"
        "--model", model,
        "--language", language,
        "--output_dir", str(output_dir),
    ]  # fmt: skip

    try:
        subprocess.run(command, check=True)
    except OSError as e:
        print_err(f"❌ Error: could not run 'whisper': {e}")
        return None
    except subprocess.CalledProcessError as e:
        print_err(
            f"❌ Error during the transcription of '{audio}' (exit {e.returncode})."
        )
        return None

    # Whisper exits with 0 even when it skips a file, so make sure the transcript is new
    if not transcript.exists() or transcript.stat().st_mtime_ns == previous_mtime:
        print_err(f"⚠️ Could not find a new .txt transcription for '{audio}'.")
        return None
    return transcript


def process_audio(
    audio: Path, model: str, language: str, formats: Sequence[str], output_root: Path
) -> bool:
    """Transcribes, formats and exports one audio file; returns True on success."""
    if not audio.is_file():
        print_err(f"❌ The audio file '{audio}' does not exist (or is not a file).")
        return False

    output_dir = output_root / audio.stem
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        print_err(f"❌ Cannot create the output folder '{output_dir}': {e}")
        return False

    transcript = run_whisper(audio, model, language, output_dir)
    if transcript is None:
        return False

    try:
        raw_text = transcript.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as e:
        print_err(f"❌ Could not read the transcription '{transcript}': {e}")
        return False

    print(f"⚙️ Creating structured version for '{audio}'...")
    paragraphs = format_text_to_paragraphs(raw_text)
    if not paragraphs:
        print_err(f"⚠️ The transcription for '{audio}' is empty; nothing to export.")
        return False

    converted_at = datetime.now().astimezone().strftime(TIMESTAMP_FORMAT)
    content = DocumentContent(audio.stem, converted_at, paragraphs)

    success = True
    for fmt in formats:
        output_file = output_dir / f"{audio.stem}.{fmt}"
        try:
            EXPORTERS[fmt].save(content, output_file)
        except Exception as e:  # the export libraries raise many different error types
            print_err(f"   ❌ Error exporting to .{fmt}: {e}")
            success = False
        else:
            print(f"   ✅ Successfully exported: {output_file}")

    try:
        readme = write_readme(output_dir, audio, converted_at, model, language)
    except OSError as e:
        print_err(f"   ❌ Error writing README.txt: {e}")
        success = False
    else:
        print(f"   📝 Conversion notes: {readme}")
    return success


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    """Parses the command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Transcribes audio using Whisper and exports to DOCX and/or PDF "
        "with proper formatting. Without --docx or --pdf, both are generated."
    )
    parser.add_argument(
        "files", nargs="+", type=Path, help="Paths to the audio files to process."
    )
    parser.add_argument(
        "--model",
        default="medium",
        help="Whisper model to use (e.g., small, medium, large). Default: medium",
    )
    parser.add_argument(
        "--language", default="es", help="Language of the audio. Default: es (Spanish)"
    )
    parser.add_argument(
        "--docx",
        action="store_true",
        help="Generate only a formatted Word (.docx) file",
    )
    parser.add_argument(
        "--pdf", action="store_true", help="Generate only a formatted PDF (.pdf) file"
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        type=Path,
        default=OUTPUT_ROOT,
        help="Base directory for the results; each audio file gets its own subfolder "
        f"(created if needed). Default: {OUTPUT_ROOT}",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    """Runs the tool and returns the process exit code."""
    args = parse_args(argv)
    # If no format flag is provided, generate every format
    formats = [fmt for fmt in EXPORTERS if getattr(args, fmt)] or list(EXPORTERS)

    print_banner()

    # Check dependencies before starting the (long) transcriptions
    missing = missing_dependencies(formats)
    if missing:
        print_err("\n❌ ABORTED: Missing dependencies. Nothing was transcribed.")
        for package, reason in missing.items():
            print_err(f"   - Missing package: {package} ({reason})")
        print_err("\n💡 Solution: Install them by running:")
        print_err(f"   pip install {' '.join(missing)}")
        print_err(
            "   or install everything using:\n   pip install -r requirements.txt\n"
        )
        return 1

    output_root: Path = args.output_dir
    try:
        output_root.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        print_err(f"❌ ABORTED: Cannot use the output directory '{output_root}': {e}")
        return 1

    files: list[Path] = args.files
    for name, count in Counter(audio.stem for audio in files).items():
        if count > 1:
            print_err(
                f"⚠️ {count} input files are named '{name}'; their outputs will collide."
            )

    failed: list[Path] = []
    try:
        for audio in files:
            if not process_audio(
                audio, args.model, args.language, formats, output_root
            ):
                failed.append(audio)
    except KeyboardInterrupt:
        print_err("\n🛑 Interrupted by the user.")
        return 130

    print(f"\n{RULE}")
    if failed:
        print_err(
            f"❌ {len(failed)} of {len(files)} audio file(s) could not be processed:"
        )
        for audio in failed:
            print_err(f"   - {audio}")
    else:
        print("🎉 All audio files have been processed successfully!")
    print(f"{RULE}\n")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
